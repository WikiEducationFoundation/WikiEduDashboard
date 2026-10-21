#!/usr/bin/env python3
"""Add a "Claim ID" column to the Fact Verification curation spreadsheet.

The spreadsheet identifies claims by article name plus the viewer's "Claim X of
N" position. claim_ids.rb (run on production) writes a CSV
mapping each (article, position) to its VerificationClaim id; this puts that id
into the spreadsheet as a new column right after "Claim #", so the rows keep
their identity as claims are removed from the pool and the positions shift.

    python3 claim_ids_merge.py fv_claim_ids.csv spreadsheet.xlsx [--out OUT] [--force]

Every claim row must resolve to exactly one id and every article's row count
must equal its highlighted-claim count in the CSV; otherwise the discrepancies
are reported and nothing is written (pass --force to write what did match).
Without --out the spreadsheet is updated in place. Requires openpyxl.
"""

import argparse
import copy
import csv
import re
import sys
from collections import Counter

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange

ARTICLE_HEADER = 'Article name'
POSITION_HEADER = 'Claim #'
ID_HEADER = 'Claim ID'


def normalize(title):
    return re.sub(r'\s+', ' ', str(title).replace('_', ' ')).strip().casefold()


def load_ids(csv_path):
    """{normalized title: {position: claim_id}}"""
    ids = {}
    with open(csv_path, newline='', encoding='utf-8') as handle:
        for row in csv.DictReader(handle):
            ids.setdefault(normalize(row['article_title']), {})[int(row['position'])] = row['claim_id']
    return ids


def resolve_title(cell_title, known_titles):
    """The CSV title a spreadsheet article name refers to, or None.

    The exercise's tiles truncate long titles with an ellipsis, and that
    truncated form is what ended up in the spreadsheet, so fall back to a
    unique prefix match.
    """
    wanted = normalize(cell_title)
    if wanted in known_titles:
        return wanted
    prefix = wanted.rstrip('.…').rstrip()
    if prefix != wanted:
        matches = [title for title in known_titles if title.startswith(prefix)]
        if len(matches) == 1:
            return matches[0]
    return None


def header_index(ws, text):
    for cell in ws[1]:
        if cell.value == text:
            return cell.column
    sys.exit(f'No "{text}" header in row 1')


def shift_validations(ws, insert_col):
    for validation in ws.data_validations.dataValidation:
        shifted = []
        for cell_range in validation.sqref.ranges:
            cell_range = CellRange(cell_range.coord)
            if cell_range.min_col >= insert_col:
                cell_range.shift(col_shift=1)
            elif cell_range.max_col >= insert_col:
                cell_range.expand(right=1)
            shifted.append(cell_range)
        validation.sqref = MultiCellRange(shifted)


def shift_column_widths(ws, insert_col):
    widths = {letter: dim.width for letter, dim in ws.column_dimensions.items() if dim.width}
    for letter in list(ws.column_dimensions):
        del ws.column_dimensions[letter]
    for letter, width in widths.items():
        index = openpyxl.utils.column_index_from_string(letter)
        if index >= insert_col:
            index += 1
        ws.column_dimensions[get_column_letter(index)].width = width


def main():
    parser = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    parser.add_argument('csv_path')
    parser.add_argument('xlsx_path')
    parser.add_argument('--out', help='write here instead of updating in place')
    parser.add_argument('--force', action='store_true',
                        help='write even when some rows could not be resolved')
    args = parser.parse_args()

    ids = load_ids(args.csv_path)
    wb = openpyxl.load_workbook(args.xlsx_path)
    ws = wb.active
    article_col = header_index(ws, ARTICLE_HEADER)
    position_col = header_index(ws, POSITION_HEADER)
    if header_index_or_none(ws) is not None:
        sys.exit(f'"{ID_HEADER}" column already present; nothing to do')
    insert_col = position_col + 1

    resolved = {}  # row -> claim id
    problems = []
    rows_per_title = Counter()
    for row in range(2, ws.max_row + 1):
        article = ws.cell(row, article_col).value
        position = ws.cell(row, position_col).value
        if article is None or position is None:
            continue  # blank row or a note row (no claim number)
        title = resolve_title(article, ids)
        if title is None:
            problems.append(f'row {row}: article "{article}" is not in the CSV')
            continue
        rows_per_title[title] += 1
        claim_id = ids[title].get(int(float(position)))
        if claim_id is None:
            problems.append(f'row {row}: "{article}" has no claim at position {position}')
            continue
        resolved[row] = claim_id

    print(f'{"article":<50} {"sheet rows":>10} {"csv claims":>10}')
    for title, positions in sorted(ids.items()):
        sheet_rows = rows_per_title.get(title, 0)
        flag = '' if sheet_rows == len(positions) else '  <-- MISMATCH'
        print(f'{title:<50} {sheet_rows:>10} {len(positions):>10}{flag}')
        if sheet_rows != len(positions):
            problems.append(f'"{title}": {sheet_rows} spreadsheet rows vs {len(positions)} claims in CSV')
    duplicates = [cid for cid, n in Counter(resolved.values()).items() if n > 1]
    if duplicates:
        problems.append(f'claim ids assigned to more than one row: {duplicates}')

    print(f'\nResolved {len(resolved)} rows to claim ids.')
    if problems:
        print(f'\n{len(problems)} problem(s):')
        for problem in problems:
            print('  ' + problem)
        if not args.force:
            sys.exit('Nothing written. Fix the discrepancies or pass --force.')

    shift_validations(ws, insert_col)
    shift_column_widths(ws, insert_col)
    ws.insert_cols(insert_col)
    header = ws.cell(1, insert_col, ID_HEADER)
    header.font = copy.copy(ws.cell(1, position_col).font)
    ws.column_dimensions[get_column_letter(insert_col)].width = 10
    for row, claim_id in resolved.items():
        ws.cell(row, insert_col, int(claim_id))

    out = args.out or args.xlsx_path
    wb.save(out)
    print(f'Wrote {out}')


def header_index_or_none(ws):
    return next((cell.column for cell in ws[1] if cell.value == ID_HEADER), None)


if __name__ == '__main__':
    main()
